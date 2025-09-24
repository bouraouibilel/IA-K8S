from docx import Document
import json

def get_heading_level(style_name):
    """Retourne le niveau du heading ou None si ce n'est pas un heading"""
    if style_name.startswith('Heading'):
        try:
            return int(style_name.replace('Heading ', ''))
        except ValueError:
            return None
    return None

def add_paragraph_to_section(section_stack, para_text):
    """Ajoute un paragraphe au dernier élément de la pile"""
    if not section_stack:
        section_stack.append({"title": None, "paragraphs": [para_text], "subsections": []})
    else:
        section_stack[-1]["paragraphs"].append(para_text)

def docx_to_json_hierarchical(path):
    doc = Document(path)
    root = {"subsections": []}
    section_stack = []

    for para in doc.paragraphs:
        heading_level = get_heading_level(para.style.name)
        if heading_level:
            # Créer une nouvelle section
            new_section = {"title": para.text, "paragraphs": [], "subsections": []}

            # Déterminer où l'ajouter dans la hiérarchie
            while len(section_stack) >= heading_level:
                section_stack.pop()

            if section_stack:
                section_stack[-1]["subsections"].append(new_section)
            else:
                root["subsections"].append(new_section)

            section_stack.append(new_section)
        elif para.text.strip():
            add_paragraph_to_section(section_stack, para.text.strip())

    return json.dumps(root["subsections"], ensure_ascii=False, indent=2)

# Exemple d'utilisation
json_output = docx_to_json_hierarchical("myDoc.docx")
with open("../PASRAU-Docs/resultat-small.json", "w", encoding="utf-8") as f: f.write(json_output)
